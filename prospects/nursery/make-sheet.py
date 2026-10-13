#!/usr/bin/env python3
"""Writes nursery-prospects.xlsx from data/*.json and grades/*.json: one row per nursery, worst grades first.
Usage: python3 prospects/nursery/make-sheet.py"""
import glob, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
recs = [r for f in sorted(glob.glob(os.path.join(HERE, 'data', '*.json'))) for r in json.load(open(f))]
grades = {g['id']: g for f in glob.glob(os.path.join(HERE, 'grades', '*.json')) for g in json.load(open(f))}
ORDER = {'Bad': 0, 'Average': 1, 'NoSite': 2, 'Good': 3, 'Excluded': 4}
FILL = {'Bad': 'F8D7DA', 'Average': 'FFF3CD', 'Good': 'D4EDDA', 'NoSite': 'E2E3E5', 'Excluded': 'F2F2F2'}

cols = [('Grade', 10), ('Why', 60), ('Nursery', 34), ('Website', 38), ('Town', 16), ('Council', 18), ('Sector', 10),
        ('Places', 8), ('Under 3s', 9), ('Lowest grade', 12), ('Grades published', 15), ('Phone', 15), ('Email', 32),
        ('Manager', 22), ('Provider', 32), ('Builder', 12), ('© year', 8), ('Funded hours mentioned', 12),
        ('Book a visit mentioned', 12), ('Care Inspectorate no.', 15), ('Id', 9)]
wb = Workbook(); ws = wb.active; ws.title = 'Nurseries'
ws.append([c for c, _ in cols])
for i, (_, w) in enumerate(cols, 1): ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]: c.font = Font(bold=True)
ws.freeze_panes = 'A2'

def key(r):
    g = grades.get(r['id'], {}).get('grade', 'NoSite')
    return (ORDER.get(g, 9), -(r.get('places') or 0))
for r in sorted(recs, key=key):
    g = grades.get(r['id'], {}); s = r.get('site') or {}; m = s.get('mentions') or {}
    ws.append([g.get('grade'), g.get('reason'), r['name'], r.get('website') or '', r['town'], r['council'], r['sector'],
               r.get('places'), 'Yes' if r.get('under3s') else 'No', r['grades'].get('min'), r['grades'].get('published'),
               r.get('phone'), r.get('email'), r.get('manager'), r['provider'], s.get('builder'), s.get('copyrightYear'),
               'Yes' if m.get('fundedHours') else '', 'Yes' if m.get('visit') else '', r['csNumber'], r['id']])
    ws.cell(ws.max_row, 1).fill = PatternFill('solid', fgColor=FILL.get(g.get('grade'), 'FFFFFF'))
ws.auto_filter.ref = ws.dimensions
out = os.path.join(HERE, 'nursery-prospects.xlsx')
wb.save(out)
print(f'{len(recs)} rows -> {out}')
